"""Parse supported document types to plain text."""
from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def parse_document(path: Path) -> tuple[str, list[tuple[int | None, str, str]]]:
    """
    Returns (full_text, segments) where each segment is (page_or_none, heading_path, text).
    For simpler chunking we concatenate segments into blocks later; segments help heading_path.
    """
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix in (".xlsx", ".xlsm"):
        return _parse_xlsx(path)
    if suffix in (".txt", ".md"):
        text = path.read_text(encoding="utf-8", errors="replace")
        return text, [(None, "", text)]
    raise ValueError(f"不支持的文件类型: {suffix}")


def _parse_pdf(path: Path) -> tuple[str, list[tuple[int | None, str, str]]]:
    reader = PdfReader(str(path))
    segments: list[tuple[int | None, str, str]] = []
    parts: list[str] = []
    for i, page in enumerate(reader.pages):
        t = page.extract_text() or ""
        parts.append(t)
        segments.append((i + 1, "", t))
    return "\n\n".join(parts), segments


def _parse_docx(path: Path) -> tuple[str, list[tuple[int | None, str, str]]]:
    doc = DocxDocument(str(path))
    lines: list[str] = []
    heading_stack: list[str] = []
    segments: list[tuple[int | None, str, str]] = []
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        style = para.style.name if para.style else ""
        if style.startswith("Heading"):
            level = 1
            try:
                level = int(style.replace("Heading", "").strip() or "1")
            except ValueError:
                level = 1
            while len(heading_stack) >= level:
                heading_stack.pop()
            heading_stack.append(text)
        hp = " / ".join(heading_stack)
        lines.append(text)
        segments.append((None, hp, text))
    full = "\n".join(lines)
    return full, segments if segments else [(None, "", full)]


def _parse_xlsx(path: Path) -> tuple[str, list[tuple[int | None, str, str]]]:
    wb = load_workbook(str(path), read_only=True, data_only=True)
    rows_out: list[str] = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows_out.append(f"## 工作表: {sheet}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            line = "\t".join(cells).strip()
            if line:
                rows_out.append(line)
    full = "\n".join(rows_out)
    return full, [(None, "", full)]
